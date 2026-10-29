import os
import shutil
import time
import subprocess
import streamlit as st
import pandas as pd
from docx import Document

# Cek apakah modul comtypes tersedia (khusus Windows)
try:
  import comtypes.client

  WINDOWS_PDF_SUPPORT = True
except ImportError:
  WINDOWS_PDF_SUPPORT = False

st.title("🎓 Otomatisasi Sertifikat Mail Merge to PDF")
st.write(
    "Upload file template Word (menggunakan tag bawaan Mail Merge seperti"
    " `«Nama»`) dan Excel data peserta."
)

if not WINDOWS_PDF_SUPPORT:
  st.warning(
      "⚠️ Peringatan: Pustaka `comtypes` belum terinstall. Jalankan `pip install"
      " comtypes` di terminal."
  )

uploaded_word = st.file_uploader("Upload Template Sertifikat (.docx)", type=["docx"])
uploaded_excel = st.file_uploader(
    "Upload Data Peserta (.xlsx)", type=["xlsx"]
)

if uploaded_word and uploaded_excel:
  df = pd.read_excel(uploaded_excel)
  st.write("Preview Data Peserta:", df.head())

  st.subheader("Pengaturan Penamaan File PDF")
  kolom_nama = st.selectbox("Pilih Kolom Nama Peserta", df.columns)
  kolom_status = st.selectbox("Pilih Kolom Status", df.columns)
  kolom_nomor = st.selectbox("Pilih Kolom Nomor Sertifikat", df.columns)

  if st.button("Generate Sertifikat ke PDF"):
    # 1. Paksa matikan sisa proses background Word agar file tidak terkunci
    if WINDOWS_PDF_SUPPORT:
      try:
        subprocess.run(
            ["taskkill", "/f", "/im", "WINWORD.EXE"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
      except Exception:
        pass

    output_dir = "output_sertifikat_pdf"
    if os.path.exists(output_dir):
      shutil.rmtree(output_dir, ignore_errors=True)
    os.makedirs(output_dir, exist_ok=True)

    progress_bar = st.progress(0)
    total_data = len(df)

    temp_template_path = "temp_template.docx"
    with open(temp_template_path, "wb") as f:
      f.write(uploaded_word.getbuffer())

    word_app = None
    if WINDOWS_PDF_SUPPORT:
      try:
        comtypes.CoInitialize()
        word_app = comtypes.client.CreateObject("Word.Application")
        word_app.Visible = False
        word_app.DisplayAlerts = 0
      except Exception as e:
        st.error(
            f"Gagal menginisialisasi Microsoft Word di background: {e}"
        )


    # Fungsi pengganti teks yang mendeteksi tag mail merge «Nama» serta mempertahankan font asli
    def replace_mailmerge_tags(paragraph, key, value):
      # Mendukung format «Nama», {{Nama}}, dan [Nama]
      tags = [f"«{key}»", f"{{{{{key}}}}}", f"[{key}]"]

      for tag in tags:
        if tag in paragraph.text:
          for run in paragraph.runs:
            if tag in run.text:
              run.text = run.text.replace(tag, str(value))
              return True
          # Jika tag terpecah di antara beberapa run, gabungkan ke run pertama
          full_text = "".join([r.text for r in paragraph.runs])
          if tag in full_text:
            new_text = full_text.replace(tag, str(value))
            paragraph.runs[0].text = new_text
            for r in paragraph.runs[1:]:
              r.text = ""
            return True
      return False


    for index, row in df.iterrows():
      # Buka template Word bersih untuk setiap baris peserta
      doc = Document(temp_template_path)

      # Ganti placeholder di paragraf
      for p in doc.paragraphs:
        for key in df.columns:
          replace_mailmerge_tags(p, key, row[key])

      # Ganti placeholder di dalam tabel (jika ada)
      for table in doc.tables:
        for row_t in table.rows:
          for cell in row_t.cells:
            for p in cell.paragraphs:
              for key in df.columns:
                replace_mailmerge_tags(p, key, row[key])

      # Ambil data untuk penamaan file PDF
      nama_peserta = str(row[kolom_nama]).strip()
      status_peserta = str(row[kolom_status]).strip()
      nomor_sertifikat = str(row[kolom_nomor]).strip()

      # Bersihkan karakter terlarang untuk nama file Windows
      nama_file = (
          f"{nama_peserta}_{status_peserta}_{nomor_sertifikat}"
          .replace("/", "_")
          .replace("\\", "_")
          .replace(":", "_")
          .replace("*", "_")
          .replace("?", "_")
          .replace('"', "_")
          .replace("<", "_")
          .replace(">", "_")
          .replace("|", "_")
      )

      docx_temp_path = os.path.abspath(
          os.path.join(output_dir, f"temp_{index}.docx")
      )
      pdf_output_path = os.path.abspath(
          os.path.join(output_dir, f"{nama_file}.pdf")
      )

      doc.save(docx_temp_path)

      # Konversi ke PDF lewat background Word
      if word_app and WINDOWS_PDF_SUPPORT:
        success = False
        for attempt in range(3):
          try:
            word_doc = word_app.Documents.Open(docx_temp_path)
            word_doc.SaveAs2(pdf_output_path, FileFormat=17)
            word_doc.Close(SaveChanges=False)
            success = True
            break
          except Exception:
            time.sleep(0.5)

        if not success:
          st.warning(
              f"⚠️ Baris {index} ({nama_peserta}) gagal diconvert otomatis."
          )

      time.sleep(0.05)
      if os.path.exists(docx_temp_path):
        try:
          os.remove(docx_temp_path)
        except PermissionError:
          pass

      progress_bar.progress((index + 1) / total_data)

    if word_app:
      try:
        word_app.Quit()
        comtypes.CoUninitialize()
      except:
        pass

    if os.path.exists(temp_template_path):
      os.remove(temp_template_path)

    st.success(
        f"Berhasil membuat {total_data} file sertifikat PDF! Mengompres ke"
        " bentuk ZIP..."
    )

    # Kompres folder menjadi file ZIP
    zip_filename = "Semua_Sertifikat_PDF"
    shutil.make_archive(zip_filename, "zip", output_dir)
    zip_path = f"{zip_filename}.zip"

    if os.path.exists(output_dir):
      shutil.rmtree(output_dir, ignore_errors=True)

    with open(zip_path, "rb") as f:
      st.download_button(
          label="📥 Download Semua Sertifikat PDF (.ZIP)",
          data=f,
          file_name="Sertifikat_Peserta_PDF.zip",
          mime="application/zip",
      )